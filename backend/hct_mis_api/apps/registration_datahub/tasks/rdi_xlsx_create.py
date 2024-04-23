import logging
from collections import defaultdict
from datetime import datetime
from functools import partial
from io import BytesIO
from typing import Any, Callable, Dict, Optional, Union

from django.contrib.gis.geos import Point
from django.core.files import File
from django.core.files.storage import default_storage
from django.db import transaction
from django.utils import timezone

import openpyxl
from django_countries.fields import Country
from hct_mis_api.apps.geo.models import Country as GeoCountry, Area
from openpyxl.cell import Cell
from openpyxl.worksheet.worksheet import Worksheet

from hct_mis_api.apps.activity_log.models import log_create
from hct_mis_api.apps.core.models import BusinessArea
from hct_mis_api.apps.core.utils import SheetImageLoader, timezone_datetime
from hct_mis_api.apps.household.models import (
    COLLECT_TYPE_FULL,
    COLLECT_TYPE_NONE,
    COLLECT_TYPE_PARTIAL,
    COLLECT_TYPE_UNKNOWN,
    HEAD,
    NON_BENEFICIARY,
    ROLE_ALTERNATE,
    ROLE_PRIMARY,
)
from hct_mis_api.apps.household.models import (
    BankAccountInfo,
    Document,
    DocumentType,
    Household,
    Individual,
    IndividualIdentity,
    IndividualRoleInHousehold,
)
from hct_mis_api.apps.registration_data.models import ImportData, RegistrationDataImport, RegistrationDataImportDatahub
from hct_mis_api.apps.registration_datahub.tasks.deduplicate import DeduplicateTask
from hct_mis_api.apps.registration_datahub.tasks.rdi_base_create import (
    RdiBaseCreateTask,
)
from hct_mis_api.apps.registration_datahub.tasks.utils import collectors_str_ids_to_list
from hct_mis_api.apps.utils.age_at_registration import calculate_age_at_registration

logger = logging.getLogger(__name__)


class RdiXlsxCreateTask(RdiBaseCreateTask):
    """
    Works on valid XLSX files, parsing them and creating households/individuals
    in the Registration Datahub. Once finished it will update the status of
    that registration data import instance.
    """

    def __init__(self) -> None:
        self.image_loader: Optional[SheetImageLoader] = None
        self.business_area = None
        self.households = {}
        self.documents = {}
        self.identities = {}
        self.household_identities = {}
        self.individuals = []
        self.collectors = defaultdict(list)
        self.bank_accounts = defaultdict(dict)

    def _handle_collect_individual_data(
        self,
        value: Any,
        header: str,
        row_num: int,
        individual: Individual,
        *args: Any,
        **kwargs: Any,
    ) -> str:
        try:
            return {
                "FULL": COLLECT_TYPE_FULL,
                "PARTIAL": COLLECT_TYPE_PARTIAL,
                "NONE": COLLECT_TYPE_NONE,
                "UNKNOWN": COLLECT_TYPE_UNKNOWN,
            }[value]
        except KeyError:
            return COLLECT_TYPE_UNKNOWN

    def _handle_bank_account_fields(
        self,
        value: Any,
        header: str,
        row_num: int,
        individual: Individual,
        *args: Any,
        **kwargs: Any,
    ) -> None:
        if value is None:
            return

        name = header.replace("_i_c", "").replace("pp_", "")

        self.bank_accounts[f"individual_{row_num}"]["individual"] = individual
        self.bank_accounts[f"individual_{row_num}"][name] = value

    def _handle_document_fields(
        self,
        value: Any,
        header: str,
        row_num: int,
        individual: Individual,
        *args: Any,
        **kwargs: Any,
    ) -> None:
        if value is None:
            return

        header = header.replace("_no", "").replace("pp_", "")
        common_header = f"individual_{row_num}_{header}"
        document_key = header.replace("_i_c", "").strip()
        document_data = self.documents.get(
            common_header,
            {
                "individual": individual,
                "key": document_key,
            },
        )

        document_data["value"] = value
        self.documents[common_header] = document_data

    def _handle_document_photo_fields(
        self,
        cell: Any,
        row_num: int,
        individual: Individual,
        header: str,
        *args: Any,
        **kwargs: Any,
    ) -> None:
        if not self.image_loader.image_in(cell.coordinate):
            return

        header = header.replace("_photo_i_c", "_i_c").replace("pp_", "")
        document_key = header.replace("_i_c", "").strip()
        common_header = f"individual_{row_num}_{header}"
        document_data = self.documents.get(
            common_header,
            {
                "individual": individual,
                "key": document_key,
            },
        )
        file = self._handle_image_field(cell)
        document_data["photo"] = file
        self.documents[common_header] = document_data

    def _handle_document_issuing_country_fields(
        self,
        value: Any,
        header: str,
        row_num: int,
        individual: Individual,
        *args: Any,
        **kwargs: Any,
    ) -> None:
        if value is None:
            return

        header = header.replace("_issuer_i_c", "_i_c").replace("pp_", "")
        document_key = header.replace("_i_c", "").strip()
        common_header = f"individual_{row_num}_{header}"
        document_data = self.documents.get(
            common_header,
            {
                "individual": individual,
                "key": document_key,
            },
        )
        document_data["issuing_country"] = Country(value)
        self.documents[common_header] = document_data

    def _handle_image_field(
        self,
        cell: Any,
        is_flex_field: bool = False,
        is_field_required: bool = False,
        *args: Any,
        **kwargs: Any,
    ) -> Union[File, str, None]:
        if self.image_loader.image_in(cell.coordinate):
            image = self.image_loader.get(cell.coordinate)
            file_name = f"{cell.coordinate}-{timezone.now()}.jpg"
            file_io = BytesIO()

            image.save(file_io, image.format)

            file = File(file_io, name=file_name)

            if is_flex_field:
                return default_storage.save(file_name, file)

            return file
        return "" if is_field_required is True else None

    def _handle_decimal_field(
        self,
        cell: Any,
        is_flex_field: bool = False,
        is_field_required: bool = False,
        *args: Any,
        **kwargs: Any,
    ) -> Any:
        value = cell.value
        if not is_flex_field:
            return value
        return float(value)

    def _handle_bool_field(
        self,
        cell: Any,
        is_flex_field: bool = False,
        is_field_required: bool = False,
        *args: Any,
        **kwargs: Any,
    ) -> Any:
        value = cell.value
        if isinstance(value, str):
            if value.lower() == "false":
                return False
            elif value.lower() == "true":
                return True
        return value

    def _handle_geopoint_field(self, value: Any, *args: Any, **kwargs: Any) -> Union[str, Point]:
        if not value:
            return ""

        values_as_list = value.split(",")
        longitude = values_as_list[0].strip()
        latitude = values_as_list[1].strip()

        return Point(x=float(longitude), y=float(latitude), srid=4326)

    def _handle_datetime(
        self,
        cell: Any,
        is_flex_field: bool = False,
        is_field_required: bool = False,
        *args: Any,
        **kwargs: Any,
    ) -> datetime:
        return timezone_datetime(cell.value)

    def _handle_identity_fields(
        self,
        value: Any,
        header: str,
        row_num: int,
        individual: Individual,
        *args: Any,
        **kwargs: Any,
    ) -> None:
        if value is None:
            return

        partner = "WFP" if "scope_id" in header else "UNHCR"

        identities_data = self.identities.get(f"individual_{row_num}_{partner}")

        if identities_data:
            identities_data["number"] = value
            identities_data["partner"] = partner
        else:
            self.identities[f"individual_{row_num}_{partner}"] = {
                "individual": individual,
                "number": value,
                "partner": partner,
            }

    def _handle_identity_photo(
        self,
        cell: Any,
        row_num: int,
        header: str,
        individual: Individual,
        *args: Any,
        **kwargs: Any,
    ) -> None:
        if not self.image_loader.image_in(cell.coordinate):
            return

        partner = "WFP" if "scope_id" in header else "UNHCR"

        identity_data = self.identities.get(f"individual_{row_num}_{partner}")

        image = self.image_loader.get(cell.coordinate)
        file_name = f"{cell.coordinate}-{timezone.now()}.jpg"
        file_io = BytesIO()

        image.save(file_io, image.format)

        file = File(file_io, name=file_name)

        if identity_data:
            identity_data["photo"] = file
        else:
            self.identities[f"individual_{row_num}_{partner}"] = {
                "individual": individual,
                "photo": file,
                "partner": partner,
            }

    def _handle_identity_issuing_country_fields(
        self,
        value: Any,
        header: str,
        row_num: int,
        individual: Individual,
        *args: Any,
        **kwargs: Any,
    ) -> None:
        if value is None:
            return

        partner = "WFP" if "scope_id" in header else "UNHCR"

        identities_data = self.identities.get(f"individual_{row_num}_{partner}")

        if identities_data:
            identities_data["issuing_country"] = Country(value)
        else:
            self.identities[f"individual_{row_num}_{partner}"] = {
                "individual": individual,
                "issuing_country": Country(value),
                "partner": partner,
            }

    def _handle_collectors(
        self,
        value: Any,
        header: str,
        individual: Individual,
        *args: Any,
        **kwargs: Any,
    ) -> None:
        list_of_ids = collectors_str_ids_to_list(value)
        if list_of_ids is None:
            return

        for hh_id in list_of_ids:
            if not hh_id:
                continue
            role = ROLE_PRIMARY if header == "primary_collector_id" else ROLE_ALTERNATE
            self.collectors[hh_id].append(IndividualRoleInHousehold(individual=individual, role=role))

    def _create_bank_accounts_infos(self) -> None:
        bank_accounts_infos_to_create = [
            BankAccountInfo(**bank_account_info) for bank_account_info in self.bank_accounts.values()
        ]

        BankAccountInfo.objects.bulk_create(bank_accounts_infos_to_create)

    def _create_documents(self) -> None:
        docs_to_create = []
        for document_data in self.documents.values():
            issuing_country = document_data.get("issuing_country")
            doc_type = DocumentType.objects.get(key=document_data["key"])
            photo = document_data.get("photo")
            individual = document_data.get("individual")
            obj = Document(
                country=issuing_country,
                document_number=document_data.get("value"),
                photo=photo,
                individual=individual,
                type=doc_type,
            )

            docs_to_create.append(obj)

        Document.objects.bulk_create(docs_to_create)

    def _create_identities(self) -> None:
        identities_to_create = [
            IndividualIdentity(
                partner=identity["partner"],
                individual=identity["individual"],
                document_number=identity["number"],
                country=identity["issuing_country"],
            )
            for identity in self.identities.values()
        ]

        IndividualIdentity.objects.bulk_create(identities_to_create)

    def _create_collectors(self) -> None:
        collectors_to_create = []
        for hh_id, collectors_list in self.collectors.items():
            for collector in collectors_list:
                collector.household_id = self.households.get(hh_id).pk
                collectors_to_create.append(collector)
        IndividualRoleInHousehold.objects.bulk_create(collectors_to_create)

    @staticmethod
    def _validate_birth_date(obj_to_create: Any) -> Any:
        birth_date = obj_to_create.birth_date

        if obj_to_create.birth_date < datetime(1923, 1, 1):
            obj_to_create.birth_date = datetime(1923, 1, 1)
        if obj_to_create.birth_date > datetime.today():
            obj_to_create.birth_date = datetime(2022, 4, 25)

        if birth_date != obj_to_create.birth_date:
            obj_to_create.estimated_birth_date = True

        return obj_to_create

    def _create_objects(
            self,
            sheet: Worksheet,
            registration_data_import_datahub: RegistrationDataImportDatahub
    ) -> None:
        complex_fields: Dict[str, Dict[str, Callable]] = {
            "individuals": {
                "photo_i_c": self._handle_image_field,
                "primary_collector_id": self._handle_collectors,
                "alternate_collector_id": self._handle_collectors,
                "pregnant_i_c": self._handle_bool_field,
                "fchild_hoh_i_c": self._handle_bool_field,
                "child_hoh_i_c": self._handle_bool_field,
                "bank_name_i_c": self._handle_bank_account_fields,
                "bank_account_number_i_c": self._handle_bank_account_fields,
                "debit_card_number_i_c": self._handle_bank_account_fields,
                "account_holder_name_i_c": self._handle_bank_account_fields,
                "bank_branch_name_i_c": self._handle_bank_account_fields,
                "first_registration_date_i_c": self._handle_datetime,
                "unhcr_id_no_i_c": self._handle_identity_fields,
                "unhcr_id_photo_i_c": self._handle_identity_photo,
                "unhcr_id_issuer_i_c": self._handle_identity_issuing_country_fields,
                "scope_id_no_i_c": self._handle_identity_fields,
                "scope_id_photo_i_c": self._handle_identity_photo,
                "scope_id_issuer_i_c": self._handle_identity_issuing_country_fields,
            },
            "households": {
                "consent_sign_h_c": self._handle_image_field,
                "hh_geopoint_h_c": self._handle_geopoint_field,
                "fchild_hoh_h_c": self._handle_bool_field,
                "child_hoh_h_c": self._handle_bool_field,
                "consent_h_c": self._handle_bool_field,
                "first_registration_date_h_c": self._handle_datetime,
                "collect_individual_data": self._handle_collect_individual_data,
            },
        }
        document_complex_types: Dict[str, Callable] = {}
        for document_type in DocumentType.objects.all():
            document_complex_types[f"{document_type.key}_i_c"] = self._handle_document_fields
            document_complex_types[f"{document_type.key}_no_i_c"] = self._handle_document_fields
            document_complex_types[f"{document_type.key}_photo_i_c"] = self._handle_document_photo_fields
            document_complex_types[f"{document_type.key}_issuer_i_c"] = self._handle_document_issuing_country_fields

        complex_fields["individuals"].update(document_complex_types)
        complex_types: Dict[str, Callable] = {
            "GEOPOINT": self._handle_geopoint_field,
            "IMAGE": self._handle_image_field,
            "DECIMAL": self._handle_decimal_field,
            "BOOL": self._handle_bool_field,
        }

        registration_data_import = RegistrationDataImport.objects.get(id=registration_data_import_datahub.hct_id)
        program_id = registration_data_import.program.id

        sheet_title = str(sheet.title.lower())
        if sheet_title == "households":
            obj = partial(
                Household,
                registration_data_import=registration_data_import,
                program_id=program_id,
                business_area=self.business_area
            )
        elif sheet_title == "individuals":
            obj = partial(
                Individual,
                registration_data_import=registration_data_import,
                program_id=program_id,
                business_area=self.business_area
            )
        else:
            raise ValueError(f"Unhandled sheet label '{sheet.title!r}'")

        first_row = sheet[1]
        households_to_update = []

        def has_value(cell: Cell) -> bool:
            if cell.value is None:
                return False
            if isinstance(cell.value, str):
                return cell.value.strip() != ""
            return True

        for row in sheet.iter_rows(min_row=3):
            if not any(has_value(cell) for cell in row):
                continue
            try:
                obj_to_create = obj()

                household_id = None

                excluded = ("age",)
                for cell, header_cell in zip(row, first_row):
                    try:
                        header = header_cell.value
                        combined_fields = self.COMBINED_FIELDS
                        current_field = combined_fields.get(header, {})

                        if not current_field and header not in complex_fields[sheet_title]:
                            continue

                        is_not_image = current_field.get("type") != "IMAGE"

                        cell_value = cell.value
                        if isinstance(cell_value, str):
                            cell_value = cell_value.strip()

                        is_not_required_and_empty = (
                            not current_field.get("required") and cell.value is None and is_not_image
                        )
                        if header in excluded:
                            continue
                        if is_not_required_and_empty:
                            continue

                        if header == "household_id":
                            temp_value = cell_value
                            if isinstance(temp_value, float) and temp_value.is_integer():
                                temp_value = int(temp_value)
                            household_id = str(temp_value)
                            if sheet_title == "individuals":
                                obj_to_create.household = self.households.get(household_id)

                        if header in complex_fields[sheet_title]:
                            fn_complex: Callable = complex_fields[sheet_title][header]
                            value = fn_complex(
                                value=cell_value,
                                cell=cell,
                                header=header,
                                row_num=cell.row,
                                individual=obj_to_create if sheet_title == "individuals" else None,
                                household=obj_to_create if sheet_title == "households" else None,
                                is_field_required=current_field.get("required", False),
                            )
                            if value is not None:
                                setattr(
                                    obj_to_create,
                                    combined_fields[header]["name"],
                                    value,
                                )
                        elif (
                            hasattr(
                                obj_to_create,
                                combined_fields[header]["name"],
                            )
                            and header != "household_id"
                        ):
                            value = self._cast_value(cell_value, header)
                            if value in (None, ""):
                                continue

                            if header == "relationship_i_c" and value == HEAD:
                                household = self.households.get(household_id)
                                if household is not None:
                                    household.head_of_household = obj_to_create
                                    households_to_update.append(household)

                            if header == "country_h_c":
                                setattr(obj_to_create, "country", GeoCountry.objects.get(iso_code3=value))
                            elif header in ("admin1_h_c", "admin2_h_c"):
                                setattr(obj_to_create, combined_fields[header]["name"], Area.objects.get(p_code=value))
                            else:
                                setattr(
                                    obj_to_create,
                                    combined_fields[header]["name"],
                                    value,
                                )
                        elif header in self.FLEX_FIELDS[sheet_title]:
                            value = self._cast_value(cell_value, header)
                            type_name = self.FLEX_FIELDS[sheet_title][header]["type"]
                            if type_name in complex_types:
                                fn_flex: Callable = complex_types[type_name]
                                value = fn_flex(
                                    value=cell_value,
                                    cell=cell,
                                    header=header,
                                    is_flex_field=True,
                                    is_field_required=current_field.get("required", False),
                                )
                            if value is not None:
                                obj_to_create.flex_fields[header] = value
                    except Exception as e:
                        raise Exception(
                            f"Error processing cell {header_cell} with `{cell}`: {e.__class__.__name__}({e})"
                        ) from e

                obj_to_create.last_registration_date = obj_to_create.first_registration_date
                obj_to_create.detail_id = row[0].row
                if sheet_title == "households":
                    obj_to_create.set_admin_areas()
                    self.households[household_id] = obj_to_create
                else:
                    if household_id is None:
                        obj_to_create.relationship = NON_BENEFICIARY
                    obj_to_create = self._validate_birth_date(obj_to_create)
                    obj_to_create.age_at_registration = calculate_age_at_registration(
                        registration_data_import_datahub, str(obj_to_create.birth_date)
                    )
                    self.individuals.append(obj_to_create)
            except Exception as e:
                raise Exception(f"Error processing row {row[0].row}: {e.__class__.__name__}({e})") from e

        print(self.households.values())
        if sheet_title == "households":
            Household.objects.bulk_create(self.households.values())
        else:
            Individual.objects.bulk_create(self.individuals)
            Household.objects.bulk_update(
                households_to_update,
                ["head_of_household"],
                1000,
            )
            self._create_documents()
            self._create_identities()
            self._create_collectors()
            self._create_bank_accounts_infos()

    @transaction.atomic(using="default")
    def execute(
        self, registration_data_import_id: str, import_data_id: str, business_area_id: str, program_id: str
    ) -> None:
        registration_data_import_datahub = RegistrationDataImportDatahub.objects.select_for_update().get(
            id=registration_data_import_id,
        )
        registration_data_import_datahub.import_done = RegistrationDataImportDatahub.STARTED
        registration_data_import_datahub.save()

        import_data = ImportData.objects.get(id=import_data_id)

        self.business_area = BusinessArea.objects.get(id=business_area_id)

        wb = openpyxl.load_workbook(import_data.file, data_only=True)

        # households objects have to be created first
        worksheets = (wb["Households"], wb["Individuals"])
        logger.info("Starting import of %s", registration_data_import_datahub.id)
        for sheet in worksheets:
            self.image_loader = SheetImageLoader(sheet)
            self._create_objects(sheet, registration_data_import_datahub)

        registration_data_import_datahub.import_done = RegistrationDataImportDatahub.DONE
        registration_data_import_datahub.save()
        old_rdi_mis = RegistrationDataImport.objects.get(id=registration_data_import_datahub.hct_id)
        if not self.business_area.postpone_deduplication:
            logger.info("Starting deduplication of %s", registration_data_import_datahub.id)
            rdi_mis = RegistrationDataImport.objects.get(id=registration_data_import_datahub.hct_id)
            rdi_mis.status = RegistrationDataImport.DEDUPLICATION
            rdi_mis.save()
            DeduplicateTask(self.business_area.slug, str(program_id)).deduplicate_imported_individuals(
                registration_data_import_datahub=registration_data_import_datahub
            )
            logger.info("Finished deduplication of %s", registration_data_import_datahub.id)
        else:
            rdi_mis = RegistrationDataImport.objects.get(id=registration_data_import_datahub.hct_id)
            rdi_mis.status = RegistrationDataImport.IN_REVIEW
            rdi_mis.save()
            log_create(
                RegistrationDataImport.ACTIVITY_LOG_MAPPING,
                "business_area",
                None,
                rdi_mis.program_id,
                old_rdi_mis,
                rdi_mis,
            )
